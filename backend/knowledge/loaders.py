"""Extract plain text from uploaded knowledge files."""
from __future__ import annotations

import io
from pathlib import Path

SUPPORTED = {".txt", ".md", ".pdf", ".docx"}


class UnsupportedFile(ValueError):
    pass


def load_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise UnsupportedFile(f"Chỉ hỗ trợ {', '.join(sorted(SUPPORTED))} — nhận được '{ext or filename}'")
    if ext in {".txt", ".md"}:
        return data.decode("utf-8", errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def is_markdown(filename: str) -> bool:
    return Path(filename).suffix.lower() == ".md"
